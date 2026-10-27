"""文件作用说明：将需求与技术 Markdown 源稿生成统一排版的可编辑 Word 文档，供评审和交付使用。"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "文档"


def font_settings(style, size, bold=False):
    """设置中英文字体、字号和黑色字形，返回已配置样式。"""
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in list(fonts.attrib):
        if "theme" in key.lower():
            del fonts.attrib[key]
    for key, value in {"ascii": "Arial", "hAnsi": "Arial", "eastAsia": "Hiragino Sans GB"}.items():
        fonts.set(qn("w:" + key), value)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn("w:" + attr), None)
    return style


def configure_document(document, title):
    """配置页面、正文、标题层级与页码，使用适合中文技术文档的简洁样式。"""
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.80)
    section.right_margin = Inches(0.80)
    section.footer_distance = Inches(0.30)
    font_settings(document.styles["Normal"], 11)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.22
    normal.space_after = Pt(6)
    normal.widow_control = True
    for name, size in (("Title", 24), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 12.5)):
        font_settings(document.styles[name], size, name != "Subtitle")
        fmt = document.styles[name].paragraph_format
        fmt.space_before = Pt(14 if name == "Heading 1" else 10)
        fmt.space_after = Pt(7)
        fmt.keep_with_next = True
        borders = document.styles[name].element.xpath("./w:pPr/w:pBdr")
        for border in borders:
            border.getparent().remove(border)
    document.styles["Title"].paragraph_format.space_before = Pt(0)
    font_settings(document.styles["Caption"], 9.5)
    for name in ("List Bullet", "List Number"):
        font_settings(document.styles[name], 11)
        document.styles[name].paragraph_format.line_spacing = 1.18
        document.styles[name].paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("第 ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)
    document.core_properties.title = title
    document.core_properties.subject = "在线客服 SaaS 网页首版"
    document.core_properties.author = "邓云川"
    document.core_properties.keywords = "SaaS, Go, PostgreSQL, 客服, AI装修, 租户模板"


def add_inline(paragraph, text):
    """解析少量加粗和行内代码标记，保留可编辑的正文内容。"""
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def cell_shading(cell, fill):
    """设置表格单元格背景色。"""
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def table_widths(headers):
    """按照各类表格的语义分配列宽，减少长说明在窄列换行。"""
    count = len(headers)
    presets = {
        ("编号", "已确定要求", "首版交付口径"): [0.65, 1.85, 4.40],
        ("编号", "场景", "通过条件"): [0.60, 1.80, 4.50],
        ("模式", "租户能力", "编辑结果"): [1.20, 3.60, 2.10],
        ("资源范围", "建议操作", "身份与规则"): [2.80, 1.70, 2.40],
        ("需求", "技术模块", "主要验收"): [1.35, 3.15, 2.40],
    }
    if tuple(headers) in presets:
        return presets[tuple(headers)]
    if count == 2:
        return [1.75, 5.15]
    if count == 3:
        if headers[0] in ("编号", "参数"):
            return [0.60, 2.05, 4.25] if headers[0] == "编号" else [1.65, 1.45, 3.80]
        if headers[0] in ("字段", "实体字段"):
            return [2.55, 1.30, 3.05]
        return [1.65, 1.80, 3.45]
    if count == 4:
        return [1.10, 1.85, 1.75, 2.20]
    if count == 5:
        return [1.70, 1.25, 1.35, 1.30, 1.30]
    return [6.9 / count] * count


def add_table(document, lines):
    """将 Markdown 表格写入 Word，设置重复表头、显式边框、留白和自适应行高。"""
    records = []
    for line in lines:
        cells = [item.strip() for item in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", item) for item in cells):
            continue
        records.append(cells)
    headers = records[0]
    widths = table_widths(headers)
    table = document.add_table(rows=0, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "85"), ("bottom", "85"), ("left", "100"), ("right", "100")):
        margin = OxmlElement("w:" + edge)
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    props.append(margins)
    for index, values in enumerate(records):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
        for column, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            if column == 0 and value == "conversation_assignments":
                value = "conversation_\nassignments"
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            paragraph.paragraph_format.keep_with_next = False
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                run.font.bold = index == 0
                run.font.color.rgb = RGBColor(255, 255, 255) if index == 0 else RGBColor(0, 0, 0)
            cell_shading(cell, "3C4548" if index == 0 else ("F4F6F6" if index % 2 == 0 else "FFFFFF"))
    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = 0.35
    gap.add_run().font.size = Pt(3)


def build_document(source):
    """读取一份源稿并生成对应 Word 文件，返回输出路径。"""
    document = Document()
    lines = source.read_text(encoding="utf-8").splitlines()
    title = next(line[2:] for line in lines if line.startswith("# "))
    configure_document(document, title)
    after_title = 0
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("<!--"):
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
            after_title = 2
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif after_title:
            paragraph = document.add_paragraph(line, style="Subtitle" if after_title == 2 else "Normal")
            if after_title == 1:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            after_title -= 1
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.20)
            paragraph.paragraph_format.first_line_indent = Inches(-0.20)
            add_inline(paragraph, line)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
            if line.startswith("https://"):
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                paragraph.paragraph_format.space_after = Pt(7)
            elif re.match(r"[TE]\d{2} ", line):
                paragraph.paragraph_format.keep_with_next = True
        index += 1
    # 参考证据按标题和整段路径一起分页，避免行号或后续路径离开所属证据。
    paragraphs = list(document.paragraphs)
    for pos, item in enumerate(paragraphs):
        if re.match(r"E\d{2} ", item.text):
            following = pos + 1
            while following < len(paragraphs) and not re.match(r"E\d{2} ", paragraphs[following].text) and not paragraphs[following].style.name.startswith("Heading"):
                following += 1
            for reference in paragraphs[pos:following]:
                reference.paragraph_format.keep_together = True
                for run in reference.runs:
                    if run.font.name == "Menlo":
                        run.text = re.sub(r"(?=:\d)", " ", re.sub(r"(:\d+)-(\d+)(?=$|,)", r"\1‑\2", run.text))
            for reference in paragraphs[pos:following-1]:
                reference.paragraph_format.keep_with_next = True
    OUTPUT.mkdir(exist_ok=True)
    version_line = next(line for line in lines if line.startswith("版本 "))
    version = re.search(r"版本 (\d+\.\d+)", version_line).group(1)
    destination = OUTPUT / (source.stem + " v" + version + ".docx")
    document.save(destination)
    return destination


if __name__ == "__main__":
    for document_source in sorted((ROOT / "docs").glob("在线客服SaaS*.md")):
        print(build_document(document_source))
